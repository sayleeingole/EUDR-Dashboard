"""Systematic exports: risk assessment report (docx), Art. 9 checklist
(docx + xlsx), audit extract (xlsx). Naming convention:
KLK_EUDR_<country>_<region|ALL>_<cycle>_<doctype>_<YYYYMMDD>.<ext>
Every export is an audited event; files land in exports/."""

import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from openpyxl import Workbook

from . import db
from .checklist import latest_signed, supplier_minimum_items

BASE_DIR = Path(__file__).resolve().parent.parent
EXPORTS_DIR = BASE_DIR / "exports"

RATING_COLORS = {"not_negligible": RGBColor(0xA3, 0x2D, 0x2D),
                 "high": RGBColor(0xA3, 0x2D, 0x2D),
                 "standard": RGBColor(0x85, 0x4F, 0x0B),
                 "negligible": RGBColor(0x3B, 0x6D, 0x11),
                 "low": RGBColor(0x3B, 0x6D, 0x11)}


def _filename(country, region, cycle_year, doctype, ext):
    EXPORTS_DIR.mkdir(exist_ok=True)
    region_tag = (region or "ALL").replace(" ", "")
    stamp = datetime.now().strftime("%Y%m%d")
    return EXPORTS_DIR / (f"KLK_EUDR_{country.replace(' ', '')}_{region_tag}_"
                          f"{cycle_year}_{doctype}_{stamp}.{ext}")


def _audit_export(conn, doctype, path, country, actor):
    db.audit(conn, "export_generated", "export", entity_id=path.name,
             country=country, after={"doctype": doctype}, actor=actor)
    conn.commit()


def risk_report_docx(conn, cycle, country, region, actor):
    """Country/region risk assessment report. DRAFT watermark unless all ten
    sections are signed."""
    cycle_id, cycle_year = cycle["id"], cycle["year"]
    country_id, country_name = country["id"], country["name"]
    region_id = region["id"] if region else None
    region_name = region["name"] if region else None

    sections = db.get_sections(conn)
    signed, missing = {}, []
    for s in sections:
        a, level = latest_signed(conn, cycle_id, country_id, region_id,
                                 s["code"])
        if a:
            signed[s["code"]] = (a, level)
        else:
            missing.append(s["code"])
    is_final = not missing

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = f"EUDR country risk assessment — {country_name}"
    if region_name:
        title += f" ({region_name})"
    doc.add_heading(title, level=0)
    p = doc.add_paragraph()
    p.add_run(f"Commodity: palm (oil palm, palm oil, PKO and derivatives)\n"
              f"Assessment cycle: {cycle_year}\n"
              f"Operator: KLK Emmerich GmbH\n"
              f"Generated: {datetime.now().strftime('%d %B %Y %H:%M')} "
              f"by {actor}\n"
              f"Status: {'COMPLETE — all sections signed' if is_final else 'DRAFT — unsigned sections: ' + ', '.join(missing)}")
    if not is_final:
        r = doc.add_paragraph().add_run("DRAFT — NOT FOR SUBMISSION")
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        r.bold = True

    doc.add_paragraph(
        "This assessment follows Article 10(2)(a)-(n) of Regulation (EU) "
        "2023/1115 (EUDR). Every evidence record cited was individually "
        "approved by the named operator; ratings were suggested by documented "
        "rules and confirmed or overridden by the operator; narratives are "
        "operator-signed conclusions. The full decision history is available "
        "in the audit trail extract.")

    for s in sections:
        doc.add_heading(f"{s['title']} — Art. {s['article_refs']}", level=1)
        entry = signed.get(s["code"])
        if entry is None:
            doc.add_paragraph("Not signed in this cycle.")
            continue
        a, level = entry
        p = doc.add_paragraph()
        run = p.add_run(f"Rating: {a['confirmed_rating'].upper()}")
        run.bold = True
        if a["confirmed_rating"] in RATING_COLORS:
            run.font.color.rgb = RATING_COLORS[a["confirmed_rating"]]
        p.add_run(f"   (signed v{a['version']} by {a['signed_by']} on "
                  f"{a['signed_at'][:10]}, assessed at {level} level)")
        if a["suggested_rating"] and \
                a["suggested_rating"] != a["confirmed_rating"]:
            doc.add_paragraph(
                f"Suggested rating was '{a['suggested_rating']}'; overridden "
                f"with reason: {a['override_reason']}")
        doc.add_paragraph(a["narrative"])

        ev_ids = json.loads(a["evidence_ids"] or "[]")
        if ev_ids:
            doc.add_heading("Evidence basis", level=2)
            for ev_id in ev_ids:
                ev = conn.execute("SELECT * FROM evidence WHERE id=?",
                                  (ev_id,)).fetchone()
                if ev is None:
                    continue
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"{ev['claim']} ")
                src = p.add_run(f"[{ev['source_name']}"
                                + (f", {ev['url']}" if ev["url"] else "")
                                + f", retrieved {ev['retrieved']}]")
                src.font.size = Pt(8.5)
                src.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    path = _filename(country_name, region_name, cycle_year, "RiskAssessment",
                     "docx")
    doc.save(path)
    _audit_export(conn, "RiskAssessment", path, country_name, actor)
    return path, is_final


def checklist_docx(conn, cycle, country, region, generation_row, items, actor):
    country_name = country["name"]
    region_name = region["name"] if region else None
    doc = Document()
    title = f"EUDR Article 9 documentation checklist — {country_name}"
    if region_name:
        title += f" ({region_name})"
    doc.add_heading(title, level=0)
    doc.add_paragraph(
        f"Commodity: palm · Cycle: {cycle['year']} · Generated: "
        f"{generation_row['at']} by {generation_row['generated_by']} · "
        f"Generation {generation_row['generation_id']}\n"
        "Each requested document lists the risk conclusion and/or legal "
        "requirement that justifies the request.")
    current_group = None
    current_alt_group = None
    for it in items:
        if it["grouping"] != current_group:
            current_group = it["grouping"]
            current_alt_group = None
            doc.add_heading(current_group or "Other", level=1)
        if it["alt_group"] != current_alt_group:
            current_alt_group = it["alt_group"]
            if current_alt_group:
                doc.add_paragraph("Any ONE of the following documents satisfies "
                                  "this requirement:").runs[0].italic = True
        p = doc.add_paragraph(style="List Bullet 2" if it["alt_group"] else "List Bullet")
        r = p.add_run(it["document_name"])
        r.bold = True
        if it["escalation"]:
            p.add_run(f"  ({it['escalation']})")
        if it["is_mandatory"] == "yes":
            p.add_run("  [Mandatory]")
        elif it["is_mandatory"] == "conditional":
            cond = f" — {it['mandatory_condition']}" if it["mandatory_condition"] else ""
            p.add_run(f"  [Conditional{cond}]")
        elif it["is_mandatory"] == "no":
            p.add_run("  [Optional]")
        if it["description"]:
            doc.add_paragraph(it["description"]).paragraph_format.left_indent \
                = Pt(18)
        just = doc.add_paragraph("Justification: "
                                 + " | ".join(json.loads(it["justification"])))
        just.paragraph_format.left_indent = Pt(18)
        for run in just.runs:
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    path = _filename(country_name, region_name, cycle["year"], "Art9Checklist",
                     "docx")
    doc.save(path)
    _audit_export(conn, "Art9Checklist", path, country_name, actor)
    return path


def _group_sort_key(name):
    """Universal, then legal-baseline areas (A1, A2, ...), then risk sections
    (S1, S2, ...) — a plain string sort would put 'Risk S10' before 'Risk S3'
    since '1' < '3' as characters."""
    name = name or ""
    if name.startswith("Universal"):
        return (0, 0, name)
    m = re.match(r"Area A(\d+)", name)
    if m:
        return (1, int(m.group(1)), name)
    m = re.match(r"Risk S(\d+)", name)
    if m:
        return (2, int(m.group(1)), name)
    return (3, 0, name)


def checklist_docx_supplier(conn, cycle, country, region, generation_row, items, actor):
    """Bare-minimum supplier version: mandatory baseline docs + genuinely
    triggered risk add-ons only, no legal citations, no internal
    justification text — just what to send and why (in one short line, only
    for risk items, where the supplier can't otherwise guess)."""
    country_name = country["name"]
    region_name = region["name"] if region else None
    supplier_items = sorted(supplier_minimum_items(items),
                            key=lambda it: _group_sort_key(it["grouping"]))

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = f"Documents needed — EUDR compliance, {country_name}"
    if region_name:
        title += f" ({region_name})"
    doc.add_heading(title, level=0)
    doc.add_paragraph(
        "To meet the requirements of the EU Deforestation Regulation (EUDR), "
        "please provide the documents listed below for your supply to KLK "
        "Emmerich GmbH. If a document does not exist for your operation or "
        "in your country, please let us know and explain why instead.")

    current_group = None
    current_alt_group = None
    for it in supplier_items:
        if it["grouping"] != current_group:
            current_group = it["grouping"]
            current_alt_group = None
            doc.add_heading(current_group or "Other", level=1)
        if it["alt_group"] != current_alt_group:
            current_alt_group = it["alt_group"]
            if current_alt_group:
                doc.add_paragraph("Please provide ONE of the "
                                  "following:").runs[0].italic = True
        p = doc.add_paragraph(style="List Bullet 2" if it["alt_group"] else "List Bullet")
        r = p.add_run(it["document_name"])
        r.bold = True
        if it.get("escalation") == "proof-level":
            p.add_run("  — please include supporting records, not just a "
                      "declaration")
        if it.get("stream") == "risk-addon" or \
                (it.get("stream") == "merged" and it.get("escalation")):
            doc.add_paragraph(
                "Requested based on your supply chain's current risk "
                "profile.").paragraph_format.left_indent = Pt(18)

    path = _filename(country_name, region_name, cycle["year"],
                     "SupplierChecklist", "docx")
    doc.save(path)
    _audit_export(conn, "SupplierChecklist", path, country_name, actor)
    return path


def checklist_xlsx(conn, cycle, country, region, generation_row, items, actor):
    country_name = country["name"]
    region_name = region["name"] if region else None
    wb = Workbook()
    ws = wb.active
    ws.title = "Art9 checklist"
    ws.append(["Group", "Document", "Stream", "Mandatory", "Alternatives",
               "Escalation", "Justification", "Generation", "Generated at",
               "Generated by"])
    alt_group_sizes = {}
    for it in items:
        if it["alt_group"]:
            alt_group_sizes[it["alt_group"]] = alt_group_sizes.get(it["alt_group"], 0) + 1
    for it in items:
        mandatory = {"yes": "Mandatory", "conditional": "Conditional", "no": "Optional"} \
            .get(it["is_mandatory"], "")
        if it["is_mandatory"] == "conditional" and it["mandatory_condition"]:
            mandatory += f" ({it['mandatory_condition']})"
        alt = (f"Any 1 of {alt_group_sizes[it['alt_group']]}"
              if it["alt_group"] else "")
        ws.append([it["grouping"], it["document_name"], it["stream"], mandatory, alt,
                   it["escalation"] or "",
                   " | ".join(json.loads(it["justification"])),
                   it["generation_id"], it["generated_at"],
                   it["generated_by"]])
    for col, width in zip("ABCDEFG", (28, 55, 18, 20, 14, 12, 80)):
        ws.column_dimensions[col].width = width
    path = _filename(country_name, region_name, cycle["year"], "Art9Checklist",
                     "xlsx")
    wb.save(path)
    _audit_export(conn, "Art9Checklist", path, country_name, actor)
    return path


def audit_xlsx(conn, cycle_year, country_name, actor):
    rows = db.audit_rows(conn, limit=100000,
                         country=None if country_name == "All"
                         else country_name)
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit trail"
    ws.append(["At", "Actor", "Event", "Entity", "Entity id", "Country",
               "Section", "Before", "After", "Reason"])
    for r in reversed(rows):
        ws.append([r["at"], r["actor"], r["event_type"], r["entity"],
                   r["entity_id"], r["country"], r["section_code"],
                   (r["before_state"] or "")[:500],
                   (r["after_state"] or "")[:500], r["reason"]])
    for col, width in zip("ABCDEFGHIJ", (18, 16, 22, 16, 12, 12, 8, 40, 40, 30)):
        ws.column_dimensions[col].width = width
    path = _filename(country_name if country_name != "All" else "ALL", None,
                     cycle_year, "AuditExtract", "xlsx")
    wb.save(path)
    _audit_export(conn, "AuditExtract", path, country_name, actor)
    return path
